"""An export goes in, the xlsx staff will read comes out.

Every test drives the command end to end with a fake Translator and reads the
written sheet back. Exports are built inline: no real export is ever used.
"""

import itertools
import json
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from chat_translator.cli import main
from chat_translator.translator import Message, Translation

HEADER = ["role", "original_content", "original_language", "translated_content"]


class FakeTranslator:
    """Answers from a script and remembers what it was asked."""

    def __init__(self, script: dict[str, Translation] | None = None) -> None:
        self.script = script or {}
        self.asked: list[str] = []
        self.contexts: dict[str, list[Message]] = {}

    def translate(self, content: str, context: Sequence[Message]) -> Translation:
        self.asked.append(content)
        self.contexts[content] = list(context)
        return self.script.get(content, Translation(original_language="xx", dutch=f"NL[{content}]"))


_ids = itertools.count(1)


def message(role: str, content: Any) -> dict[str, Any]:
    return {"id": f"msg-{next(_ids)}", "role": role, "content": content}


def write_export(path: Path, *conversations: list[dict[str, Any]]) -> Path:
    export = {"conversations": [{"id": f"c{i}", "messages": msgs} for i, msgs in enumerate(conversations)]}
    path.write_text(json.dumps(export, ensure_ascii=False), encoding="utf-8")
    return path


def read_sheet(path: Path) -> list[list[Any]]:
    sheet = load_workbook(path).active
    return [[cell if cell is not None else "" for cell in row] for row in sheet.iter_rows(values_only=True)]


def test_each_message_becomes_a_row_under_the_four_column_header(tmp_path: Path) -> None:
    export = write_export(
        tmp_path / "export.json",
        [message("user", "Cześć"), message("assistant", "Dzień dobry")],
    )
    translator = FakeTranslator(
        {
            "Cześć": Translation(original_language="pl", dutch="Hoi"),
            "Dzień dobry": Translation(original_language="pl", dutch="Goedendag"),
        }
    )

    main([str(export)], translator=translator)

    assert read_sheet(tmp_path / "export.xlsx") == [
        HEADER,
        ["user", "Cześć", "pl", "Hoi"],
        ["assistant", "Dzień dobry", "pl", "Goedendag"],
    ]


def test_text_that_looks_like_a_formula_stays_text(tmp_path: Path) -> None:
    export = write_export(tmp_path / "export.json", [message("user", "=) dzięki")])
    translator = FakeTranslator({"=) dzięki": Translation(original_language="pl", dutch="=) bedankt")})

    main([str(export)], translator=translator)

    sheet = load_workbook(tmp_path / "export.xlsx").active
    assert [(cell.value, cell.data_type) for cell in sheet[2]] == [
        ("user", "s"),
        ("=) dzięki", "s"),
        ("pl", "s"),
        ("=) bedankt", "s"),
    ]


def test_the_output_path_can_be_chosen(tmp_path: Path) -> None:
    export = write_export(tmp_path / "export.json", [message("user", "Hello")])
    chosen = tmp_path / "elsewhere" / "sheet.xlsx"
    chosen.parent.mkdir()

    main([str(export), "-o", str(chosen)], translator=FakeTranslator())

    assert read_sheet(chosen)[1][:2] == ["user", "Hello"]
    assert not (tmp_path / "export.xlsx").exists()


def test_an_empty_message_keeps_its_row_but_is_never_translated(tmp_path: Path) -> None:
    export = write_export(
        tmp_path / "export.json",
        [message("user", "Hello"), message("assistant", "")],
        [message("user", "   \n"), message("assistant", None), message("user", "Thanks")],
    )
    translator = FakeTranslator(
        {
            "Hello": Translation(original_language="en", dutch="Hallo"),
            "Thanks": Translation(original_language="en", dutch="Bedankt"),
        }
    )

    main([str(export)], translator=translator)

    assert translator.asked == ["Hello", "Thanks"]
    assert read_sheet(tmp_path / "export.xlsx") == [
        HEADER,
        ["user", "Hello", "en", "Hallo"],
        ["assistant", "", "", ""],
        ["user", "   \n", "", ""],
        ["assistant", "", "", ""],
        ["user", "Thanks", "en", "Bedankt"],
    ]


def test_a_message_is_translated_with_the_three_before_it_in_its_own_conversation(tmp_path: Path) -> None:
    export = write_export(
        tmp_path / "export.json",
        [
            message("user", "one"),
            message("assistant", "two"),
            message("user", ""),
            message("assistant", "three"),
            message("user", "four"),
            message("assistant", "five"),
        ],
        [message("user", "other conversation")],
    )
    translator = FakeTranslator()

    main([str(export)], translator=translator)

    assert translator.contexts["one"] == []
    assert translator.contexts["two"] == [Message("user", "one")]
    assert translator.contexts["five"] == [
        Message("assistant", "two"),
        Message("assistant", "three"),
        Message("user", "four"),
    ]
    assert translator.contexts["other conversation"] == []


def test_without_a_running_model_it_stops_before_any_work(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    export = write_export(tmp_path / "export.json", [message("user", "Hello")])
    monkeypatch.setenv("OLLAMA_HOST", "http://127.0.0.1:9")  # nothing listens here

    exit_code = main([str(export)])

    assert exit_code != 0
    assert "Ollama" in capsys.readouterr().err
    assert not (tmp_path / "export.xlsx").exists()


def test_a_model_on_another_machine_is_refused_before_any_work(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    export = write_export(tmp_path / "export.json", [message("user", "Hello")])
    monkeypatch.setenv("OLLAMA_HOST", "https://ollama.example.com")

    exit_code = main([str(export)])

    assert exit_code != 0
    assert "not on this machine" in capsys.readouterr().err
    assert not (tmp_path / "export.xlsx").exists()
