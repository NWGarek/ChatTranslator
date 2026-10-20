"""The real local model on the made-up sample. Slow, so skipped unless run with `-m ollama`.

This is the evidence that the model labels languages correctly and keeps form
links intact; the pipeline itself is covered with a fake in test_convert.py.
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from chat_translator.cli import main

pytestmark = pytest.mark.ollama

SAMPLE = Path(__file__).parent / "fixtures" / "made-up-sample.json"
URL = re.compile(r"https?://\S+")

# The original language of every message in the sample, conversation by
# conversation. The fourth flexwerker starts in English and switches to Romanian.
EXPECTED_LANGUAGES = ["pl"] * 6 + ["uk"] * 5 + ["nl"] * 4 + ["en", "en", "ro", "ro", "ro"]


def test_the_sample_is_labelled_and_translated_with_its_links_intact(tmp_path: Path) -> None:
    sheet_path = tmp_path / "sample.xlsx"

    assert main([str(SAMPLE), "-o", str(sheet_path)]) == 0

    rows = list(load_workbook(sheet_path).active.iter_rows(min_row=2, values_only=True))
    assert [row[2] for row in rows] == EXPECTED_LANGUAGES
    for _, original, _, dutch in rows:
        assert dutch, "every message in the sample has a translation"
        assert URL.findall(original) == URL.findall(dutch)


def test_a_model_that_is_not_pulled_stops_the_run_with_how_to_fix_it(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = main([str(SAMPLE), "-o", str(tmp_path / "sample.xlsx"), "--model", "no-such-model:latest"])

    assert exit_code != 0
    assert "ollama pull no-such-model:latest" in capsys.readouterr().err
    assert not (tmp_path / "sample.xlsx").exists()
